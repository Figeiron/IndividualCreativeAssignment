import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from core.exporter import Exporter
from core.response import Response, TextBox, TableBox, PlotBox

class ExcelExporter(Exporter):
    display_name = "Експорт в Excel"
    BOLD_FONT = Font(bold=True)
    CENTER_ALIGNMENT = Alignment(horizontal='center', vertical='center')
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    def __init__(self, context):
        super().__init__(context)
        self._handlers = {
            TextBox: self._handle_textbox,
            TableBox: self._handle_tablebox,
            PlotBox: self._handle_plotbox
        }

    def export(self, destination: str, data: Response = None):
        data_to_export = data or self.last_response
        
        if not data_to_export:
            return Response(boxes=[TextBox(text="Немає даних для експорту. Спочатку виконайте якусь команду.")])

        if not destination.endswith(".xlsx"):
            destination += ".xlsx"

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Exported Data"

            current_row = 1
            for box in data_to_export.boxes:
                handler = self._handlers.get(type(box))
                if handler:
                    current_row = handler(ws, box, current_row)

            self._adjust_column_widths(ws)

            wb.save(destination)
            return Response(boxes=[TextBox(text=f"Дані успішно експортовано у файл: {os.path.abspath(destination)}")])
        except Exception as e:
            return Response(boxes=[TextBox(text=f"Помилка при експорті: {str(e)}")])

    def _handle_textbox(self, ws, box: TextBox, current_row: int) -> int:
        lines = box.text.split('\n')
        for line in lines:
            ws.cell(row=current_row, column=1, value=line)
            current_row += 1
        return current_row + 1

    def _handle_tablebox(self, ws, box: TableBox, current_row: int) -> int:
        if not box.cells:
            return current_row

        max_r = 0
        for cell in box.cells:
            excel_row = current_row + cell.pos[0]
            excel_col = 1 + cell.pos[1]
            c = ws.cell(row=excel_row, column=excel_col, value=cell.text)
            c.border = self.THIN_BORDER
            if cell.pos[0] == 0:
                c.font = self.BOLD_FONT
                c.alignment = self.CENTER_ALIGNMENT
            max_r = max(max_r, cell.pos[0])

        return current_row + max_r + 2

    def _handle_plotbox(self, ws, box: PlotBox, current_row: int) -> int:
        if not box.plot_points:
            return current_row

        ws.cell(row=current_row, column=1, value="Дані графіку:").font = self.BOLD_FONT
        current_row += 1
        ws.cell(row=current_row, column=1, value="X").font = self.BOLD_FONT
        ws.cell(row=current_row, column=2, value="Y").font = self.BOLD_FONT

        data_start_row = current_row + 1
        for x, y in box.plot_points:
            current_row += 1
            ws.cell(row=current_row, column=1, value=x)
            ws.cell(row=current_row, column=2, value=y)
        data_end_row = current_row

        chart = ScatterChart()
        chart.title = "Графік"
        chart.style = 13
        chart.x_axis.title = 'X'
        chart.y_axis.title = 'Y'

        xvalues = Reference(ws, min_col=1, min_row=data_start_row, max_row=data_end_row)
        yvalues = Reference(ws, min_col=2, min_row=data_start_row, max_row=data_end_row)

        series = Series(yvalues, xvalues, title_from_data=False)
        chart.series.append(series)

        ws.add_chart(chart, f"D{data_start_row}")
        return current_row + 1

    def _adjust_column_widths(self, ws):
        for column_cells in ws.columns:
            length = max(len(str(cell.value)) for cell in column_cells if cell.value is not None)
            ws.column_dimensions[column_cells[0].column_letter].width = length + 2
